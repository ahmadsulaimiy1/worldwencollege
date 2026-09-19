#!/usr/bin/env python3
"""
DOES THE WORKBOOK STILL AGREE WITH THE PUBLISHED PLAN?

The PDF and the workbook are two renderings of one engine, which is the
structural guarantee that they cannot disagree. It is not a complete
guarantee: the workbook is a separate program in a separate language,
and it recomputes some quantities in Excel formulas rather than reading
them back. Those formulas are where drift lives — the Core Plan sheet
carried `revenue - delivery - acquisition - fixed` for a while after the
engine had begun charging refunds and institutional development too, and
it would have quietly told a Board a different answer from the document
sitting beside it.

This reads the workbook back and checks the figures against the engine
that produced the plan. Run by tests/masterplan.test.mjs.
"""
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "publication" / "WEC-LC Financial Model.xlsx"


def engine(expr, mod="masterplan"):
    src = (
        "import * as M from './scripts/publication/%s.mjs';"
        "process.stdout.write(JSON.stringify(%s));" % (mod, expr)
    )
    r = subprocess.run(["node", "--input-type=module", "-e", src],
                       cwd=ROOT, capture_output=True, text=True)
    if r.returncode:
        raise SystemExit("engine failed: " + r.stderr[:400])
    return json.loads(r.stdout)


def main():
    if not BOOK.exists():
        raise SystemExit("workbook not built: %s" % BOOK)
    wb = load_workbook(BOOK)
    core = engine("M.scenarios()", mod="projection")["core"]
    ws = wb["19 Core Plan"]

    failures = []

    def near(a, b, tol=2.0):
        return abs(float(a) - float(b)) <= tol

    for i, y in enumerate(core["years"]):
        rr = 7 + i
        for col, key in [(2, "calendar"), (3, "newLearners"), (4, "activeLearners"),
                         (5, "instructors"), (6, "revenue"), (7, "refunds"),
                         (9, "delivery"), (10, "acquisition"), (11, "fixed"),
                         (12, "development")]:
            got = ws.cell(rr, col).value
            if got is None or not near(got, y[key]):
                failures.append("Core Plan row %d %s: workbook %r, engine %r"
                                % (rr, key, got, y[key]))
        # The surplus cell is a FORMULA, so check the arithmetic it encodes
        # rather than a cached value openpyxl will not compute.
        formula = ws.cell(rr, 13).value
        expected = "=H%d-I%d-J%d-K%d-L%d" % (rr, rr, rr, rr, rr)
        if formula != expected:
            failures.append("Core Plan row %d surplus formula: %r, expected %r"
                            % (rr, formula, expected))
        # ...and that the arithmetic it encodes is the engine's own.
        net = y["revenue"] - y["refunds"]
        derived = net - y["delivery"] - y["acquisition"] - y["fixed"] - y["development"]
        if not near(derived, y["surplus"]):
            failures.append("row %d: the formula's arithmetic gives %.2f, engine says %.2f"
                            % (rr, derived, y["surplus"]))

    # The frontier sheet must carry the price the plan actually proposes.
    fr = wb["22 Teaching frontier"]
    proposed = engine("M.PROPOSED.committed.directed", mod="pricing")
    prices = [fr.cell(r, 1).value for r in range(7, fr.max_row + 1)]
    if proposed not in prices:
        failures.append("Teaching frontier does not contain the proposed price %r" % proposed)

    # And the evidence sheet must carry every observation the plan cites.
    ev = json.loads((ROOT / "data" / "market-evidence.json").read_text())
    want = sum(1 for k in ("gcc", "west_africa", "uk_europe", "executive_and_corporate")
               for o in ev[k]["observations"] if o.get("price_local") or o.get("price_usd"))
    mev = wb["21 Market evidence"]
    got = sum(1 for r in range(7, mev.max_row + 1)
              if mev.cell(r, 2).value and mev.cell(r, 9).value)
    if got != want:
        failures.append("Market evidence sheet has %d observations, the file has %d" % (got, want))

    for f in failures:
        print("FAIL " + f)
    print("workbook-sync: %d checks failed" % len(failures))
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
